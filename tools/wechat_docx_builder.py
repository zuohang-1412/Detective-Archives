from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FONT = os.environ.get("WECHAT_DOCX_FONT", "Microsoft YaHei")

NAVY = "17324D"
BLUE = "2E5C7A"
BRASS = "B58A4C"
INK = "26323A"
MUTED = "65727C"
PALE_BLUE = "EEF3F6"
PALE_GOLD = "F7F1E6"


@dataclass(frozen=True)
class ArticleConfig:
    source: Path
    output: Path
    eyebrow: str
    cover_alt: str
    subject: str
    keywords: str
    meta_line: str
    case_heading_pattern: str | None = None
    metadata_pattern: str | None = None
    highlight_pattern: str | None = None


def set_run_font(run, *, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_style(style, *, size, color=INK, bold=False, before=0, after=0, line=1.25):
    style.font.name = FONT
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.line_spacing = line
    paragraph.widow_control = True


def set_paragraph_shading(paragraph, fill):
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, *, side, color, size=10, space=4):
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def add_hyperlink(paragraph, label, url):
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([fonts, color, underline])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_picture_alt_text(inline_shape, description):
    properties = inline_shape._inline.docPr
    properties.set("title", description)
    properties.set("descr", description)


def add_inline_markdown(paragraph, text, *, size=11, color=INK):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=NAVY, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def parse_frontmatter(raw):
    lines = raw.splitlines()
    if not lines or lines[0].strip() != "---":
        return {}, lines
    metadata = {}
    end = None
    for index, line in enumerate(lines[1:], 1):
        if line.strip() == "---":
            end = index
            break
        if ":" in line:
            key, value = line.split(":", 1)
            metadata[key.strip()] = value.strip()
    if end is None:
        raise ValueError("Markdown front matter is not closed.")
    return metadata, lines[end + 1 :]


def article_body(lines):
    """Drop editorial notes before the article's first level-one heading."""
    for index, line in enumerate(lines):
        if line.strip().startswith("# "):
            return lines[index:]
    return lines


def resolve_asset(source, relative_path):
    asset = (source.parent / relative_path).resolve()
    if not asset.is_file():
        raise FileNotFoundError(f"Missing article asset: {asset}")
    return asset


def add_cover(doc, config, metadata):
    cover = resolve_asset(config.source, metadata["coverImage"])
    mark = resolve_asset(config.source, metadata["projectMark"])

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(9)
    picture = paragraph.add_run().add_picture(str(cover), width=Inches(6.5))
    set_picture_alt_text(picture, config.cover_alt)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    picture = paragraph.add_run().add_picture(str(mark), width=Inches(0.62))
    set_picture_alt_text(picture, "侦探档案馆项目标志")

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(config.eyebrow)
    set_run_font(run, size=9, color=BRASS, bold=True)

    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(metadata["title"])
    set_run_font(run, size=20, color=NAVY, bold=True)

    paragraph = doc.add_paragraph(style="Subtitle")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.right_indent = Inches(0.35)
    run = paragraph.add_run(metadata["summary"])
    set_run_font(run, size=10.5, color=MUTED)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(config.meta_line.format(**metadata))
    set_run_font(run, size=9, color=BLUE, bold=True)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    set_paragraph_border(rule, side="bottom", color=BRASS, size=9, space=1)


def add_heading(doc, text, level, case_heading_pattern=None):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    add_inline_markdown(paragraph, text, size=13 if level == 2 else 12, color=NAVY)
    if level == 2:
        set_paragraph_border(paragraph, side="bottom", color=BRASS, size=7, space=3)
    elif level == 3 and case_heading_pattern and re.match(case_heading_pattern, text):
        set_paragraph_shading(paragraph, PALE_BLUE)
        set_paragraph_border(paragraph, side="left", color=BRASS, size=18, space=5)
        paragraph.paragraph_format.left_indent = Inches(0.08)
        paragraph.paragraph_format.right_indent = Inches(0.04)
    return paragraph


def add_body(doc, text, metadata_pattern=None, highlight_pattern=None):
    paragraph = doc.add_paragraph(style="Normal")
    add_inline_markdown(paragraph, text)
    if metadata_pattern and re.match(metadata_pattern, text):
        paragraph.style = doc.styles["Article Metadata"]
        set_paragraph_shading(paragraph, PALE_GOLD)
    elif highlight_pattern and re.match(highlight_pattern, text):
        set_paragraph_shading(paragraph, PALE_GOLD)
        paragraph.paragraph_format.left_indent = Inches(0.08)
        paragraph.paragraph_format.right_indent = Inches(0.05)
    return paragraph


def add_quote(doc, lines):
    paragraph = doc.add_paragraph(style="Article Callout")
    set_paragraph_shading(paragraph, PALE_GOLD)
    set_paragraph_border(paragraph, side="left", color=BRASS, size=20, space=6)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        add_inline_markdown(paragraph, line, size=10.5)
    return paragraph


def add_list_item(doc, text, style):
    paragraph = doc.add_paragraph(style=style)
    add_inline_markdown(paragraph, text)
    return paragraph


def add_image(doc, config, alt, relative_path):
    image_path = resolve_asset(config.source, relative_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    picture = paragraph.add_run().add_picture(str(image_path), width=Inches(6.5))
    set_picture_alt_text(picture, alt)
    paragraph.paragraph_format.keep_with_next = True

    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(alt)
    set_run_font(run, size=8.5, color=MUTED)


def add_rule(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    set_paragraph_border(paragraph, side="bottom", color=BRASS, size=7, space=1)


def add_source_link(doc, url):
    paragraph = doc.add_paragraph(style="Source Link")
    domain = urlparse(url).netloc.removeprefix("www.")
    add_hyperlink(paragraph, f"打开来源页面（{domain}）", url)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    configure_style(styles["Normal"], size=11, after=6, line=1.25)
    configure_style(styles["Title"], size=20, color=NAVY, bold=True, after=7, line=1.1)
    configure_style(styles["Subtitle"], size=10.5, color=MUTED, after=7, line=1.2)
    configure_style(styles["Heading 1"], size=16, color=BLUE, bold=True, before=18, after=10)
    configure_style(styles["Heading 2"], size=13, color=BLUE, bold=True, before=14, after=7)
    configure_style(styles["Heading 3"], size=12, color=NAVY, bold=True, before=10, after=5)
    configure_style(styles["List Bullet"], size=11, after=4, line=1.25)
    configure_style(styles["List Number"], size=11, after=4, line=1.25)
    for style_name in ("List Bullet", "List Number"):
        paragraph = styles[style_name].paragraph_format
        paragraph.left_indent = Inches(0.375)
        paragraph.first_line_indent = Inches(-0.188)
    configure_style(styles["Caption"], size=8.5, color=MUTED, after=9, line=1.15)

    metadata = styles.add_style("Article Metadata", WD_STYLE_TYPE.PARAGRAPH)
    configure_style(metadata, size=10.5, after=3, line=1.15)
    metadata.paragraph_format.left_indent = Inches(0.08)
    metadata.paragraph_format.right_indent = Inches(0.05)
    metadata.paragraph_format.keep_together = True

    callout = styles.add_style("Article Callout", WD_STYLE_TYPE.PARAGRAPH)
    configure_style(callout, size=10.5, after=9, line=1.22)
    callout.paragraph_format.left_indent = Inches(0.20)
    callout.paragraph_format.right_indent = Inches(0.10)
    callout.paragraph_format.space_before = Pt(4)
    callout.paragraph_format.keep_together = True

    source = styles.add_style("Source Link", WD_STYLE_TYPE.PARAGRAPH)
    configure_style(source, size=9.5, color=BLUE, after=6, line=1.15)
    source.paragraph_format.left_indent = Inches(0.15)


def build_article(config):
    if not config.source.is_file():
        raise FileNotFoundError(f"Missing article source: {config.source}")

    raw = config.source.read_text(encoding="utf-8")
    metadata, lines = parse_frontmatter(raw)
    missing = {"title", "summary", "coverImage", "projectMark"} - metadata.keys()
    if missing:
        raise ValueError(f"Missing front matter fields in {config.source}: {sorted(missing)}")

    doc = Document()
    configure_document(doc)
    properties = doc.core_properties
    properties.title = metadata["title"]
    properties.subject = config.subject
    properties.author = "侦探档案馆"
    properties.keywords = config.keywords
    properties.comments = "公众号导入版：单栏、内嵌图片、无运行页眉页脚。"

    add_cover(doc, config, metadata)

    skipped_title = False
    quote_lines = []
    project_mark = metadata.get("projectMark")
    for raw_line in article_body(lines):
        line = raw_line.strip()
        if quote_lines and not line.startswith(">"):
            add_quote(doc, quote_lines)
            quote_lines = []
        if not line:
            continue
        if line == "---":
            add_rule(doc)
        elif line.startswith("# "):
            if not skipped_title:
                skipped_title = True
            else:
                add_heading(doc, line[2:], 1, config.case_heading_pattern)
        elif line.startswith("## "):
            add_heading(doc, line[3:], 2, config.case_heading_pattern)
        elif line.startswith("### "):
            add_heading(doc, line[4:], 3, config.case_heading_pattern)
        elif line.startswith(">"):
            quote_lines.append(line[1:].lstrip())
        elif match := re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line):
            if match.group(2) != project_mark:
                add_image(doc, config, match.group(1), match.group(2))
        elif line.startswith("- "):
            add_list_item(doc, line[2:], "List Bullet")
        elif re.match(r"^\d+\.\s", line):
            add_list_item(doc, re.sub(r"^\d+\.\s", "", line), "List Number")
        elif re.fullmatch(r"https?://\S+", line):
            add_source_link(doc, line)
        else:
            add_body(doc, line, config.metadata_pattern, config.highlight_pattern)
    if quote_lines:
        add_quote(doc, quote_lines)

    config.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(config.output)
    return config.output
