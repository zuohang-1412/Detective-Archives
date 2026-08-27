from __future__ import annotations

import argparse
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from build_gloria_scott_wechat_docx import build as build_gloria
from build_opening_wechat_docx import build as build_opening
from build_sherlock_wechat_docx import build as build_sherlock
from wechat_docx_builder import ROOT


OUTPUT_NAMES = {
    "opening": "侦探档案馆开馆说明-公众号导入版.docx",
    "sherlock": "夏洛克·福尔摩斯人物总档案-公众号导入版.docx",
    "gloria": "福尔摩斯档案001-格洛里亚·斯科特号-公众号导入版.docx",
}


def verify_docx(path):
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"DOCX was not created: {path}")
    with ZipFile(path) as package:
        corrupt_member = package.testzip()
        media = [name for name in package.namelist() if name.startswith("word/media/")]
    if corrupt_member:
        raise RuntimeError(f"Corrupt DOCX member in {path}: {corrupt_member}")
    if not media:
        raise RuntimeError(f"DOCX has no embedded images: {path}")

    document = Document(path)
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if not document.core_properties.title:
        raise RuntimeError(f"DOCX has no title metadata: {path}")
    if "发布配置（不复制到正文）" in body_text:
        raise RuntimeError(f"Editorial-only notes leaked into DOCX: {path}")
    if len(document.inline_shapes) != len(media):
        raise RuntimeError(f"Embedded-image count mismatch in {path}")
    descriptions = [
        element.get("descr")
        for element in document.element.body.iter(qn("wp:docPr"))
    ]
    if len(descriptions) != len(media) or not all(descriptions):
        raise RuntimeError(f"An embedded image is missing alt text: {path}")
    if not any(paragraph.style.name.startswith("Heading ") for paragraph in document.paragraphs):
        raise RuntimeError(f"DOCX has no semantic headings: {path}")

    section = document.sections[0]
    dimensions = (
        round(section.page_width.inches, 2),
        round(section.page_height.inches, 2),
        round(section.left_margin.inches, 2),
        round(section.right_margin.inches, 2),
        round(section.top_margin.inches, 2),
        round(section.bottom_margin.inches, 2),
    )
    if dimensions != (8.5, 11.0, 1.0, 1.0, 1.0, 1.0):
        raise RuntimeError(f"Unexpected page geometry in {path}: {dimensions}")
    return len(media)


def main():
    parser = argparse.ArgumentParser(description="生成当前全部微信公众号导入版 DOCX")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=ROOT / "content" / "wechat" / "publish",
        help="统一输出目录；默认 content/wechat/publish",
    )
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    outputs = [
        build_opening(output_dir / OUTPUT_NAMES["opening"]),
        build_sherlock(output_dir / OUTPUT_NAMES["sherlock"]),
        build_gloria(output_dir / OUTPUT_NAMES["gloria"]),
    ]
    for output in outputs:
        image_count = verify_docx(output)
        print(f"OK  {output}  ({image_count} embedded images)")


if __name__ == "__main__":
    main()
